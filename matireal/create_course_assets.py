import json
from pathlib import Path
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.units import cm
from pptx import Presentation
from pptx.util import Pt

# Paths
BASE = Path(__file__).parent
JSON_FILE = BASE / "course_digital_literacy_fixed.json"
MD_FILE = BASE / "course_digital_literacy_masterpiece.md"
OUT_PDF = BASE / "course_digital_literacy_syllabus.pdf"
OUT_PPTX = BASE / "course_digital_literacy_master_slides.pptx"


def load_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def load_md(path):
    if not path.exists():
        return ""
    return path.read_text(encoding="utf-8")


def make_syllabus_pdf(course, md_text, out_path):
    c = canvas.Canvas(str(out_path), pagesize=A4)
    width, height = A4
    margin = 2 * cm
    x = margin
    y = height - margin

    c.setFont("Helvetica-Bold", 18)
    c.drawString(x, y, course.get("course", {}).get("name", "Course"))
    y -= 1.2 * cm

    c.setFont("Helvetica", 10)
    meta_lines = [
        f"Code: {course['course'].get('code','')}",
        f"Start: {course['course'].get('start_date','')}",
        f"Sessions: {course['course'].get('total_sessions','')}",
        f"Hours: {course['course'].get('total_hours','')}",
        f"Density: {course.get('course',{}).get('density','medium')}",
    ]
    for line in meta_lines:
        c.drawString(x, y, line)
        y -= 0.6 * cm

    y -= 0.4 * cm
    # Short description (from md or json)
    desc = course["course"].get("description") or (
        md_text.split("\n")[0] if md_text else ""
    )
    c.setFont("Helvetica", 11)
    for chunk in split_text(desc, 90):
        c.drawString(x, y, chunk)
        y -= 0.5 * cm

    y -= 0.6 * cm
    c.setFont("Helvetica-Bold", 12)
    c.drawString(x, y, "Sessions (Overview)")
    y -= 0.8 * cm

    c.setFont("Helvetica", 10)
    for s in course.get("sessions", []):
        title = f"{s.get('id','')} - {s.get('title','')}"
        c.drawString(x, y, title)
        y -= 0.5 * cm
        objs = s.get("objectives", [])
        if isinstance(objs, list):
            for obj in objs:
                for line in split_text("- " + obj, 80):
                    c.drawString(x + 0.6 * cm, y, line)
                    y -= 0.45 * cm
        else:
            for line in split_text(str(objs), 80):
                c.drawString(x + 0.6 * cm, y, line)
                y -= 0.45 * cm
        y -= 0.3 * cm
        if y < 4 * cm:
            c.showPage()
            y = height - margin

    c.showPage()
    c.save()


def split_text(text, width):
    words = text.split()
    lines = []
    cur = ""
    for w in words:
        if len(cur) + len(w) + 1 > width:
            lines.append(cur)
            cur = w
        else:
            cur = (cur + " " + w).strip()
    if cur:
        lines.append(cur)
    return lines


def make_session_pdfs(course, out_dir):
    out_dir.mkdir(parents=True, exist_ok=True)
    for s in course.get("sessions", []):
        sid = s.get("id") or s.get("temp_id") or "s?"
        fname = out_dir / f"session_{sid}_outline.pdf"
        c = canvas.Canvas(str(fname), pagesize=A4)
        width, height = A4
        margin = 2 * cm
        x = margin
        y = height - margin
        c.setFont("Helvetica-Bold", 16)
        c.drawString(x, y, s.get("title", "Session"))
        y -= 1 * cm
        c.setFont("Helvetica-Bold", 12)
        c.drawString(x, y, "Description:")
        y -= 0.6 * cm
        c.setFont("Helvetica", 10)
        for line in split_text(s.get("description", ""), 90):
            c.drawString(x, y, line)
            y -= 0.45 * cm
            if y < 4 * cm:
                c.showPage()
                y = height - margin
        y -= 0.4 * cm
        c.setFont("Helvetica-Bold", 12)
        c.drawString(x, y, "Objectives:")
        y -= 0.6 * cm
        c.setFont("Helvetica", 10)
        objs = s.get("objectives", [])
        if isinstance(objs, list):
            for obj in objs:
                for line in split_text("- " + obj, 90):
                    c.drawString(x + 0.6 * cm, y, line)
                    y -= 0.45 * cm
                    if y < 4 * cm:
                        c.showPage()
                        y = height - margin
        else:
            for line in split_text(str(objs), 90):
                c.drawString(x + 0.6 * cm, y, line)
                y -= 0.45 * cm
        c.showPage()
        c.save()


def make_pptx(course, out_path):
    prs = Presentation()
    # Title slide
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = course["course"].get("name", "Course")
    subtitle.text = f"{course['course'].get('code','')} — Masterpiece syllabus"

    # Session slides
    for s in course.get("sessions", []):
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = s.get("title", "")
        body = slide.shapes.placeholders[1].text_frame
        body.text = ""
        objs = s.get("objectives", [])
        if isinstance(objs, list):
            for i, obj in enumerate(objs):
                p = body.add_paragraph() if i > 0 else body.paragraphs[0]
                p.text = f"- {obj}"
                p.level = 0
        else:
            body.text = str(objs)

    # Project slide
    if course.get("assignments"):
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = "Final Project & Rubric"
        body = slide.shapes.placeholders[1].text_frame
        proj = course["assignments"][0]
        body.text = proj.get("title", "Project")
        for k, v in proj.get("rubric", {}).items():
            p = body.add_paragraph()
            p.text = f"{k}: {v}"
            p.level = 1

    prs.save(str(out_path))


if __name__ == "__main__":
    course = load_json(JSON_FILE)
    md = load_md(MD_FILE)

    print("Generating syllabus PDF...")
    make_syllabus_pdf(course, md, OUT_PDF)
    print(f"Wrote: {OUT_PDF}")

    print("Generating per-session PDFs...")
    make_session_pdfs(course, BASE / "session_pdfs_digital")
    print("Wrote session PDFs in folder: session_pdfs_digital")

    print("Generating PPTX...")
    make_pptx(course, OUT_PPTX)
    print(f"Wrote: {OUT_PPTX}")
